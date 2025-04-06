import io
import os
import shutil

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from shared.docx_service import DocumentStructure, DocxService


def test_docx_with_inline_shape_without_text_frame():
    # azfunctions\tests\test_data\1.Hammam Nasayrah.docx
    test_file = "tests/test_data/1.Hammam Nasayrah.docx"
    with open(test_file, "rb") as file:
        content = file.read()

    # Process the document
    result = DocxService.get_text_from_docx(content)

    # Verify the result
    assert isinstance(result, DocumentStructure)
    assert result.text is not None
    assert result.pages is not None
    assert result.paragraphs is not None
    assert result.tables is not None
    assert result.styles is not None
    assert result.headers is not None
    assert result.footers is not None
