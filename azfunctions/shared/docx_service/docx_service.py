import io
from typing import Dict, List, Optional

from docx import Document

from .models import (
    DocumentPage,
    DocumentStructure,
    DocumentStyle,
    FontStyle,
    Line,
    TableCell,
)


class DocxService:
    @staticmethod
    def get_text_from_docx(document_content) -> DocumentStructure:
        doc = Document(io.BytesIO(document_content))
        full_text = []
        paragraphs = []
        tables = []
        styles = {}
        headers = []
        footers = []

        # Extract styles
        for style in doc.styles:
            if hasattr(style, "name") and style.name:
                font = style.font if hasattr(style, "font") else None
                styles[style.name] = DocumentStyle(
                    name=style.name,
                    font=FontStyle(
                        name=font.name if font and hasattr(font, "name") else None,
                        size=float(font.size.pt)
                        if font and hasattr(font, "size") and font.size
                        else None,
                        bold=font.bold if font and hasattr(font, "bold") else None,
                        italic=font.italic
                        if font and hasattr(font, "italic")
                        else None,
                        underline=font.underline
                        if font and hasattr(font, "underline")
                        else None,
                    )
                    if font
                    else None,
                )

        # Extract paragraphs and text
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
                paragraphs.append(para.text)

        # Extract tables
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    row_data.append(TableCell(text=cell.text))
                    full_text.append(cell.text)
                table_data.append(row_data)
            tables.append(table_data)

        # Extract text from shapes
        for shape in doc.inline_shapes:
            if hasattr(shape, "text_frame") and shape.text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    if paragraph.text.strip():
                        full_text.append(paragraph.text)
                        paragraphs.append(paragraph.text)

        if hasattr(doc, "shapes"):
            for shape in doc.shapes:
                if hasattr(shape, "text_frame") and shape.text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            full_text.append(paragraph.text)
                            paragraphs.append(paragraph.text)

        # Extract headers and footers
        for section in doc.sections:
            if section.header:
                for paragraph in section.header.paragraphs:
                    if paragraph.text.strip():
                        headers.append(paragraph.text)
                        full_text.append(paragraph.text)

            if section.footer:
                for paragraph in section.footer.paragraphs:
                    if paragraph.text.strip():
                        footers.append(paragraph.text)
                        full_text.append(paragraph.text)

        # Create document structure
        return DocumentStructure(
            text="\n".join(full_text),
            pages=[
                DocumentPage(
                    page_number=1,
                    content="\n".join(full_text),
                    lines=[Line(content=text) for text in full_text],
                    tables=tables,
                )
            ],
            paragraphs=paragraphs,
            tables=tables,
            styles=styles,
            headers=headers,
            footers=footers,
        )
